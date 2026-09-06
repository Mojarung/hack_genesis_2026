"""PPTX из отрендеренных слайдов: каждый слайд — картинка 2560×1440 на весь кадр 16:9 (нередактируемый).

    uv run --with python-pptx python presentation/tools/make_pptx.py            # дека M3 → MISIS_MOJARUNG.pptx
    uv run --with python-pptx python presentation/tools/make_pptx.py poster     # постер → MISIS_MOJARUNG_poster.pptx
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

VARIANTS = {
    "m3": ("render", "MISIS_MOJARUNG.pptx"),
    "poster": ("render-poster", "MISIS_MOJARUNG_poster.pptx"),
}

ROOT = Path(__file__).resolve().parents[1]
variant = sys.argv[1] if len(sys.argv) > 1 else "m3"
render_dir, target_name = VARIANTS[variant]
pngs = sorted((ROOT / render_dir / "png").glob("slide-*.png"))
if not pngs:
    raise SystemExit(f"нет presentation/{render_dir}/png/slide-*.png — сначала отрендерить слайды")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
for png in pngs:
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)

target = ROOT / target_name
prs.save(target)
print(target.name, len(pngs), "slides,", target.stat().st_size // 1024, "KB")
